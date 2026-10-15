#!/usr/bin/env python3
"""
Export the values saved in the Excel workbook to data/spreads.json.

This is the "Excel is the engine" route: leave the workbook open on the
Bloomberg machine so the BDH formulas keep calculating, save it every few
minutes (see README, option B), and run this script after each save. It reads
the cached values that Excel stored, so it does not need Bloomberg itself.

It also builds the starter data file the dashboard ships with:

  python export_from_workbook.py "Tema ETFs - Intraday Spreads Dashboard.xlsx" --sample

Workbook layout it expects (same as your file):
  - "Intraday Dashboard" sheet: tickers in A11:A23 and the LMM in column J
  - one "<TICKER> Intraday" sheet per ETF: bucket times in A4:A81, spreads in
    B4:B81, fund name in B1
Times in the workbook are Excel local time and are treated as US Eastern.
"""
from __future__ import annotations

import argparse
import logging
import os
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from publish import ET, UTC, git_publish, write_json_atomic

log = logging.getLogger("export")


def export(xlsx: Path) -> dict:
    wb = load_workbook(xlsx, data_only=True)

    lmm: dict[str, str] = {}
    order: list[str] = []
    dash = wb["Intraday Dashboard"]
    for r in range(11, dash.max_row + 1):
        tk = dash.cell(r, 1).value
        if isinstance(tk, str) and tk.strip():
            tk = tk.strip()
            order.append(tk)
            lmm[tk] = str(dash.cell(r, 10).value or "")

    etfs: dict[str, dict] = {}
    for ws in wb.worksheets:
        if not ws.title.endswith(" Intraday"):
            continue
        tk = ws.title.split()[0]
        series = []
        for r in range(4, 82):
            t, v = ws.cell(r, 1).value, ws.cell(r, 2).value
            if not isinstance(t, datetime) or isinstance(v, bool) or not isinstance(v, (int, float)):
                continue  # skips blanks and #N/A
            ts = t.replace(tzinfo=ET).astimezone(UTC)
            series.append([ts.strftime("%Y-%m-%dT%H:%M:%SZ"), round(float(v), 4)])
        name = ws["B1"].value
        etfs[tk] = {"name": name if isinstance(name, str) else "", "lmm": lmm.get(tk, ""), "series": series}

    order = [t for t in order if t in etfs] + [t for t in etfs if t not in order]
    last = max((e["series"][-1][0] for e in etfs.values() if e["series"]), default=None)
    if last is None:
        raise SystemExit("No cached values found. Open the workbook in Excel with Bloomberg connected, "
                         "let it calculate, save, then run this again.")
    return {
        "sessionDate": last[:10],
        "bucketMinutes": 5,
        "order": order,
        "etfs": etfs,
        "_last_observation": last,
    }


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", type=Path)
    p.add_argument("--out", type=Path, default=Path("data/spreads.json"))
    p.add_argument("--sample", action="store_true",
                   help="mark the file as sample data (the dashboard shows a notice)")
    p.add_argument("--repo-dir", help="also commit and push to this clone of the data branch")
    p.add_argument("--branch", default="data")
    p.add_argument("--keep-history", action="store_true")
    args = p.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    payload = export(args.xlsx)
    last = payload.pop("_last_observation")
    if args.sample:
        payload["sample"] = True
        payload["updated"] = last
    else:
        # When the workbook was last saved is when these numbers were current.
        saved = datetime.fromtimestamp(os.path.getmtime(args.xlsx), tz=UTC)
        payload["updated"] = saved.strftime("%Y-%m-%dT%H:%M:%SZ")

    out = Path(args.repo_dir) / "data" / "spreads.json" if args.repo_dir else args.out
    write_json_atomic(out, payload)
    log.info("Wrote %s (%d ETFs).", out, len(payload["etfs"]))
    if args.repo_dir:
        git_publish(Path(args.repo_dir), os.path.relpath(out, args.repo_dir), args.branch,
                    f"spreads {payload['updated']}", args.keep_history)


if __name__ == "__main__":
    main()
